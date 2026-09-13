"""Generate the public static API from Microsoft's official catalog workbook."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

LEVELS = {
    "End to end": "end_to_end",
    "Process area": "process_area",
    "Process": "business_process",
    "Scenario": "scenario",
    "System process": "system_process",
    "Test case": "test_case",
}
DEPTHS = {level: index for index, level in enumerate(LEVELS.values(), 1)}
PUBLISHED_ON = "2026-07-07"
VERSION = "2026-07"
SOURCE_ABOUT = "https://learn.microsoft.com/en-us/dynamics365/guidance/business-processes/about"
SOURCE_DOWNLOAD = "https://aka.ms/BusinessProcessCatalog"


def text(value: object, maximum: int = 300) -> str:
    return str(value or "").strip()[:maximum]


def values(value: object) -> list[str]:
    return sorted({
        item.strip()
        for item in re.split(r"[\n;,]+", text(value, 4000))
        if item.strip()
    })


def write(path: Path, payload: dict) -> str:
    encoded = (json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n").encode()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(encoded)
    return hashlib.sha256(encoded).hexdigest()


def build(workbook: Path, output: Path) -> dict:
    workbook_bytes = workbook.read_bytes()
    sheet = load_workbook(workbook, read_only=True, data_only=True).active
    headers = list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))
    columns = {name: index for index, name in enumerate(headers)}
    ancestors: dict[int, str] = {}
    entries = []
    counts = Counter()

    for row in sheet.iter_rows(min_row=3, values_only=True):
        source_type = text(row[columns["Work item type"]], 80)
        level = LEVELS.get(source_type)
        if not level:
            continue
        depth = DEPTHS[level]
        microsoft_id = text(row[columns["Microsoft ID"]], 100)
        sequence_id = text(row[columns["Process sequence ID"]], 80)
        title = next(
            (
                text(row[columns[f"Title {index}"]])
                for index in range(7, 0, -1)
                if text(row[columns[f"Title {index}"]])
            ),
            "",
        )
        if not microsoft_id or not sequence_id or not title:
            raise ValueError(f"Missing catalog identity: {source_type} {sequence_id}")
        parent = ancestors.get(depth - 1)
        if depth > 1 and not parent:
            raise ValueError(f"Missing ordered parent for {sequence_id}")
        entry = {
            "id": microsoft_id,
            "parent_id": parent,
            "sequence_id": sequence_id,
            "level": level,
            "title": title,
            "catalog_status": text(row[columns["Catalog status"]], 80),
            "article_status": text(row[columns["Article status"]], 80),
            "application_families": values(row[columns["Application family"]]),
            "products": values(row[columns["Products"]]),
            "module": text(row[columns["Module"]], 200),
            "microsoft_references": values(row[columns["Microsoft references"]]),
        }
        entries.append(entry)
        ancestors[depth] = microsoft_id
        for child_depth in range(depth + 1, 7):
            ancestors.pop(child_depth, None)
        counts[level] += 1

    ids = [entry["id"] for entry in entries]
    if len(ids) != len(set(ids)):
        raise ValueError("Microsoft IDs must be unique.")
    known = set(ids)
    if any(entry["parent_id"] not in known for entry in entries if entry["parent_id"]):
        raise ValueError("Catalog contains an unknown parent.")

    base = {
        "schema": "d365.business-process-api/1",
        "version": VERSION,
        "published_on": PUBLISHED_ON,
        "source": {
            "about_url": SOURCE_ABOUT,
            "download_url": SOURCE_DOWNLOAD,
            "workbook": workbook.name,
            "sha256": hashlib.sha256(workbook_bytes).hexdigest(),
        },
        "counts": dict(sorted(counts.items())),
    }
    processes = {**base, "entries": entries}
    roots = {**base, "entries": [entry for entry in entries if entry["level"] == "end_to_end"]}

    product_entries: dict[str, list[str]] = defaultdict(list)
    for entry in entries:
        for product in entry["products"]:
            product_entries[product].append(entry["id"])
    products = {
        **base,
        "products": [
            {"name": name, "count": len(process_ids), "process_ids": process_ids}
            for name, process_ids in sorted(product_entries.items())
        ],
    }

    process_sha = write(output / "processes.json", processes)
    roots_sha = write(output / "roots.json", roots)
    products_sha = write(output / "products.json", products)
    catalog = {
        **base,
        "endpoints": {
            "roots": "roots.json",
            "processes": "processes.json",
            "products": "products.json",
        },
        "sha256": {
            "roots.json": roots_sha,
            "processes.json": process_sha,
            "products.json": products_sha,
        },
    }
    write(output / "catalog.json", catalog)
    return catalog


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parent.parent / "docs" / "api" / "v1",
    )
    args = parser.parse_args()
    result = build(args.workbook, args.output)
    print(json.dumps({"version": result["version"], "counts": result["counts"]}))
