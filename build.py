#!/usr/bin/env python3
"""The one idempotent build step for the RAPP/1 static API."""
from __future__ import annotations

import argparse
import datetime
import hashlib
import json
import re
import shutil
import tempfile
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
MANIFEST = json.loads((ROOT / "manifest.json").read_text(encoding="utf-8"))
RAW = (
    f"https://raw.githubusercontent.com/{MANIFEST['owner']}/"
    f"{MANIFEST['repository']}/{MANIFEST['branch']}"
)
PAGES = f"https://{MANIFEST['owner']}.github.io/{MANIFEST['repository']}"
NOW = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
LEVELS = {
    "End to end": "end_to_end",
    "Process area": "process_area",
    "Process": "business_process",
    "Scenario": "scenario",
    "System process": "system_process",
    "Test case": "test_case",
}
DEPTHS = {level: index for index, level in enumerate(LEVELS.values(), 1)}


def text(value: object, maximum: int = 300) -> str:
    return str(value or "").strip()[:maximum]


def values(value: object) -> list[str]:
    return sorted({
        item.strip()
        for item in re.split(r"[\n;,]+", text(value, 4000))
        if item.strip()
    })


def stable_json(path: Path, payload: dict, timestamp_keys=("generated",)) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        try:
            old = json.loads(path.read_text(encoding="utf-8"))
            candidate = dict(payload)
            for key in timestamp_keys:
                if key in old:
                    candidate[key] = old[key]
            if candidate == old:
                payload = candidate
        except (OSError, ValueError, json.JSONDecodeError):
            pass
    encoded = (json.dumps(payload, ensure_ascii=False, indent=2) + "\n").encode()
    if not path.exists() or path.read_bytes() != encoded:
        path.write_bytes(encoded)
    return hashlib.sha256(encoded).hexdigest()


def workbook_path(explicit: Path | None) -> tuple[Path, bool]:
    if explicit:
        return explicit, False
    handle = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    handle.close()
    target = Path(handle.name)
    request = urllib.request.Request(
        MANIFEST["source"]["workbook_url"],
        headers={"User-Agent": "rapp-static-api-builder/1.0"},
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        target.write_bytes(response.read(20 * 1024 * 1024))
    return target, True


def extract(workbook: Path) -> tuple[list[dict], dict]:
    sheet = load_workbook(workbook, read_only=True, data_only=True).active
    headers = list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))
    columns = {name: index for index, name in enumerate(headers)}
    ancestors: dict[int, str] = {}
    entries = []
    counts = Counter()

    for row in sheet.iter_rows(min_row=3, values_only=True):
        source_type = text(row[columns["Work item type"]], 80)
        level = LEVELS.get(source_type)
        if not level:
            continue
        depth = DEPTHS[level]
        microsoft_id = text(row[columns["Microsoft ID"]], 100)
        sequence_id = text(row[columns["Process sequence ID"]], 80)
        title = next(
            (
                text(row[columns[f"Title {index}"]])
                for index in range(7, 0, -1)
                if text(row[columns[f"Title {index}"]])
            ),
            "",
        )
        if not microsoft_id or not sequence_id or not title:
            raise ValueError(f"Missing catalog identity: {source_type} {sequence_id}")
        parent = ancestors.get(depth - 1)
        if depth > 1 and not parent:
            raise ValueError(f"Missing ordered parent for {sequence_id}")
        entries.append({
            "id": microsoft_id,
            "parent_id": parent,
            "sequence_id": sequence_id,
            "level": level,
            "title": title,
            "catalog_status": text(row[columns["Catalog status"]], 80),
            "article_status": text(row[columns["Article status"]], 80),
            "application_families": values(row[columns["Application family"]]),
            "products": values(row[columns["Products"]]),
            "module": text(row[columns["Module"]], 200),
            "microsoft_references": values(row[columns["Microsoft references"]]),
        })
        ancestors[depth] = microsoft_id
        for child_depth in range(depth + 1, 7):
            ancestors.pop(child_depth, None)
        counts[level] += 1

    ids = [entry["id"] for entry in entries]
    if len(ids) != len(set(ids)):
        raise ValueError("Microsoft IDs must be unique.")
    known = set(ids)
    if any(entry["parent_id"] not in known for entry in entries if entry["parent_id"]):
        raise ValueError("Catalog contains an unknown parent.")
    return entries, dict(sorted(counts.items()))


def build(explicit_workbook: Path | None = None) -> dict:
    workbook, temporary = workbook_path(explicit_workbook)
    try:
        workbook_bytes = workbook.read_bytes()
        entries, counts = extract(workbook)
    finally:
        if temporary:
            workbook.unlink(missing_ok=True)

    base = {
        "schema": "rapp-d365-business-processes/1.0",
        "name": MANIFEST["name"],
        "version": MANIFEST["version"],
        "generated": NOW,
        "published_on": MANIFEST["published_on"],
        "source": {
            "about_url": MANIFEST["source"]["about_url"],
            "download_url": MANIFEST["source"]["download_url"],
            "workbook": MANIFEST["source"]["workbook_name"],
            "sha256": hashlib.sha256(workbook_bytes).hexdigest(),
        },
        "counts": counts,
    }
    api = ROOT / "api" / "v1"
    processes = {**base, "entries": entries}
    roots = {**base, "entries": [entry for entry in entries if entry["level"] == "end_to_end"]}
    product_entries: dict[str, list[str]] = defaultdict(list)
    for entry in entries:
        for product in entry["products"]:
            product_entries[product].append(entry["id"])
    products = {
        **base,
        "products": [
            {"name": name, "count": len(process_ids), "process_ids": process_ids}
            for name, process_ids in sorted(product_entries.items())
        ],
    }

    hashes = {
        "processes.json": stable_json(api / "processes.json", processes),
        "roots.json": stable_json(api / "roots.json", roots),
        "products.json": stable_json(api / "products.json", products),
    }
    versions = ROOT / "versions"
    for name, digest in hashes.items():
        source = api / name
        target = versions / name.removesuffix(".json") / f"{digest[:12]}.json"
        if not target.exists():
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(source, target)

    registry = {
        "schema": "rapp-static-api/1.0",
        "name": MANIFEST["name"],
        "title": MANIFEST["title"],
        "description": MANIFEST["description"],
        "spec": "https://raw.githubusercontent.com/kody-w/rapp-static-apis/main/SPEC.md",
        "raw_base": RAW,
        "pages_base": PAGES,
        "generated": NOW,
        "summary": {"entries": len(entries), **counts},
        "source": base["source"],
        "entries": [
            {
                "name": name.removesuffix(".json"),
                "schema": "rapp-d365-business-processes/1.0",
                "url": f"{RAW}/api/v1/{name}",
                "pages_url": f"{PAGES}/api/v1/{name}",
                "sha256": digest,
                "sha8": digest[:12],
                "version_url": f"{RAW}/versions/{name.removesuffix('.json')}/{digest[:12]}.json",
            }
            for name, digest in hashes.items()
        ],
    }
    stable_json(ROOT / "registry.json", registry)
    stable_json(api / "status.json", {
        "schema": "rapp-d365-business-processes-status/1.0",
        "name": MANIFEST["name"],
        "generated": NOW,
        "ok": True,
        "version": MANIFEST["version"],
        "entries": len(entries),
        "source_sha256": base["source"]["sha256"],
    })
    stable_json(api / "badge.json", {
        "schemaVersion": 1,
        "label": "business processes",
        "message": str(len(entries)),
        "color": "0078d4",
    }, timestamp_keys=())
    stable_json(ROOT / ".well-known" / "mcp.json", {
        "schema": "rapp-d365-business-processes-mcp/1.0",
        "name": MANIFEST["name"],
        "protocolVersion": "2024-11-05",
        "generated": NOW,
        "resources": [
            {"uri": entry["url"], "name": entry["name"], "mimeType": "application/json"}
            for entry in registry["entries"]
        ],
    })
    stable_json(ROOT / ".well-known" / "agent-protocol.json", {
        "schema": "rapp-d365-business-processes-agent-protocol/1.0",
        "name": MANIFEST["name"],
        "generated": NOW,
        "actions": [
            {"name": "catalog", "method": "GET", "url": f"{RAW}/registry.json", "auth": "none"},
            {"name": "roots", "method": "GET", "url": f"{RAW}/api/v1/roots.json", "auth": "none"},
            {"name": "processes", "method": "GET", "url": f"{RAW}/api/v1/processes.json", "auth": "none"},
            {"name": "products", "method": "GET", "url": f"{RAW}/api/v1/products.json", "auth": "none"},
        ],
    })
    (ROOT / ".nojekyll").touch()
    llms = (
        "# Dynamics 365 Business Process API\n\n"
        "> RAPP/1 read-only static API generated from Microsoft's official catalog.\n\n"
        f"- Registry: {RAW}/registry.json\n"
        f"- Status: {RAW}/api/v1/status.json\n"
        f"- Roots: {RAW}/api/v1/roots.json\n"
        f"- Processes: {RAW}/api/v1/processes.json\n"
        f"- Products: {RAW}/api/v1/products.json\n"
    )
    (ROOT / "llms.txt").write_text(llms, encoding="utf-8")
    return registry


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path)
    args = parser.parse_args()
    result = build(args.workbook)
    print(json.dumps(result["summary"]))
